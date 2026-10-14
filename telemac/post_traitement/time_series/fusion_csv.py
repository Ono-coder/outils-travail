# -*- coding: utf-8 -*-
"""
Fusion des CSV intermédiaires (extraction TELEMAC par sous-domaine) en un
fichier Excel unique, avec identification des points de contrôle.

- Lit la liste de points (fichier de time series TELEMAC : x  y  id  nom)
- Pour chaque point configuré, recherche la colonne (nœud X..._Y...) la plus
  proche, à moins de DIST_MAX mètres.
- En-tête du fichier fusionné :
      ligne 1 : X du nœud retenu
      ligne 2 : Y du nœud retenu
      ligne 3 : nom de la variable (nom du point)
- Les points sans nœud à moins de DIST_MAX sont signalés (console + onglet
  "Correspondance") et apparaissent en colonne vide marquée "MANQUANT".

Usage :
    python fusion_csv.py [fichier_points] [dossier_csv] [sortie.xlsx] [--dist-max 100]
    python fusion_csv.py --config config.toml
    python fusion_csv.py --config config.toml --sortie autre.xlsx --dist-max 50
Les arguments en ligne de commande restent prioritaires sur le fichier de
configuration, lui-même prioritaire sur les valeurs par défaut.
"""
import argparse
import csv
import glob
import math
import os
import re
import sys
from collections import defaultdict

if sys.version_info >= (3, 11):
    import tomllib
else:
    import tomli as tomllib

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
except ImportError:
    print("❌ Le module openpyxl est requis : pip install openpyxl")
    sys.exit(1)

# --------------------------------------------------------------------------- #
#  Paramètres par défaut                                                      #
# --------------------------------------------------------------------------- #
POINTS_FILE  = "time_series.txt"
INPUT_FOLDER = "csv_intermediaires"
OUTPUT_XLSX  = "fusion_time_series.xlsx"
DIST_MAX     = 100.0   # m : distance max entre point configuré et nœud retenu

COL_RE = re.compile(r"^X(-?[\d.]+)_Y(-?[\d.]+)$")


def to_float(s):
    return float(str(s).strip().replace(",", "."))


# --------------------------------------------------------------------------- #
#  Lecture de la liste de points                                              #
# --------------------------------------------------------------------------- #
def lire_points(path):
    """Retourne une liste de dicts {x, y, id, nom} dans l'ordre du fichier.
    Les lignes d'en-tête TELEMAC (ex. '1 51', '0 1000000 2') sont ignorées :
    seules les lignes 'x y id nom' (≥ 4 champs) sont retenues."""
    points = []
    with open(path, "r", encoding="utf-8", errors="replace") as f:
        for n, line in enumerate(f, 1):
            parts = line.split()
            if len(parts) < 4:
                continue
            try:
                x, y = to_float(parts[0]), to_float(parts[1])
            except ValueError:
                continue
            pid = parts[2]
            nom = " ".join(parts[3:])
            points.append({"x": x, "y": y, "id": pid, "nom": nom, "ligne": n})
    return points


# --------------------------------------------------------------------------- #
#  Lecture et fusion des CSV                                                  #
# --------------------------------------------------------------------------- #
def lire_csv(input_folder):
    csv_files = sorted(glob.glob(os.path.join(input_folder, "*.csv")))
    print(f"📁 {len(csv_files)} fichiers CSV trouvés dans {input_folder}")

    all_data = defaultdict(dict)   # {temps(str): {colonne: valeur(float)}}
    all_columns = set()

    for i, file_path in enumerate(csv_files, 1):
        print(f"📖 [{i}/{len(csv_files)}] {os.path.basename(file_path)}")
        try:
            with open(file_path, "r", encoding="utf-8") as f:
                reader = csv.DictReader(f, delimiter="\t")
                for row in reader:
                    temps = row["temps"]
                    for col, val in row.items():
                        if col == "temps" or col is None or val is None:
                            continue
                        val = val.strip()
                        if not val:
                            continue
                        try:
                            all_data[temps][col] = to_float(val)
                        except ValueError:
                            all_data[temps][col] = val
                        all_columns.add(col)
        except Exception as e:
            print(f"  ⚠ Erreur: {e}")
            continue

    print(f"📊 {len(all_data)} pas de temps, {len(all_columns)} colonnes (nœuds)")
    return all_data, all_columns


def coords_colonnes(all_columns):
    """{colonne: (x, y)} pour les colonnes au format X..._Y..."""
    coords = {}
    for col in all_columns:
        m = COL_RE.match(col.strip())
        if m:
            coords[col] = (float(m.group(1)), float(m.group(2)))
        else:
            print(f"  ⚠ Colonne au format inattendu ignorée : {col}")
    return coords


# --------------------------------------------------------------------------- #
#  Association points configurés ↔ nœuds                                      #
# --------------------------------------------------------------------------- #
def associer(points, coords, dist_max):
    for p in points:
        best_col, best_d = None, math.inf
        for col, (cx, cy) in coords.items():
            d = math.hypot(cx - p["x"], cy - p["y"])
            if d < best_d:
                best_col, best_d = col, d
        p["dist"] = best_d if best_col else None
        if best_col is not None and best_d <= dist_max:
            p["col"] = best_col
            p["xn"], p["yn"] = coords[best_col]
            p["statut"] = "OK"
        else:
            p["col"] = None
            p["xn"] = p["yn"] = None
            p["statut"] = "MANQUANT"
            p["col_proche"] = best_col

    # Nœud partagé par plusieurs points
    usage = defaultdict(list)
    for p in points:
        if p["col"]:
            usage[p["col"]].append(p["nom"])
    for p in points:
        if p["col"] and len(usage[p["col"]]) > 1:
            p["statut"] = "OK (nœud partagé : " + ", ".join(usage[p["col"]]) + ")"

    non_associees = sorted(set(coords) - set(usage))
    return non_associees


# --------------------------------------------------------------------------- #
#  Export Excel                                                               #
# --------------------------------------------------------------------------- #
def exporter_xlsx(output_xlsx, points, all_data, coords, non_associees):
    wb = Workbook(write_only=True)
    bold = Font(bold=True)
    red = Font(bold=True, color="C00000")
    fill_miss = PatternFill("solid", fgColor="F8CBAD")

    def cell(ws, value, font=None, fill=None):
        from openpyxl.cell import WriteOnlyCell
        c = WriteOnlyCell(ws, value=value)
        if font:
            c.font = font
        if fill:
            c.fill = fill
        return c

    # ---- Feuille données ------------------------------------------------- #
    ws = wb.create_sheet("Donnees")
    ws.freeze_panes = "B4"

    # colonnes : points configurés dans l'ordre, puis nœuds non associés
    colonnes = [(p["nom"], p["col"], p["xn"] if p["col"] else p["x"],
                 p["yn"] if p["col"] else p["y"], p["col"] is None) for p in points]
    for col in non_associees:
        x, y = coords[col]
        colonnes.append(("NON_ASSOCIE", col, x, y, False))

    ws.append([cell(ws, "X", bold)] +
              [cell(ws, c[2], red if c[4] else bold, fill_miss if c[4] else None) for c in colonnes])
    ws.append([cell(ws, "Y", bold)] +
              [cell(ws, c[3], red if c[4] else bold, fill_miss if c[4] else None) for c in colonnes])
    ws.append([cell(ws, "temps (s)", bold)] +
              [cell(ws, (c[0] + " - MANQUANT") if c[4] else c[0],
                    red if c[4] else bold, fill_miss if c[4] else None) for c in colonnes])

    sorted_times = sorted(all_data.keys(), key=to_float)
    for temps in sorted_times:
        d = all_data[temps]
        ws.append([to_float(temps)] + [d.get(c[1]) if c[1] else None for c in colonnes])

    # ---- Feuille correspondance ------------------------------------------ #
    ws2 = wb.create_sheet("Correspondance")
    ws2.append([cell(ws2, h, bold) for h in
                ["id", "Nom", "X configuré", "Y configuré", "X nœud", "Y nœud",
                 "Distance (m)", "Statut", "Nœud le plus proche (si manquant)"]])
    for p in points:
        miss = p["col"] is None
        f = red if miss else None
        ws2.append([
            cell(ws2, p["id"], f), cell(ws2, p["nom"], f),
            cell(ws2, p["x"], f), cell(ws2, p["y"], f),
            cell(ws2, p["xn"], f), cell(ws2, p["yn"], f),
            cell(ws2, round(p["dist"], 2) if p["dist"] is not None else None, f),
            cell(ws2, p["statut"], f),
            cell(ws2, p.get("col_proche") if miss else None, f),
        ])
    if non_associees:
        ws2.append([])
        ws2.append([cell(ws2, "Nœuds présents dans les CSV mais associés à aucun point :", bold)])
        for col in non_associees:
            ws2.append([None, col, None, None, coords[col][0], coords[col][1]])

    wb.save(output_xlsx)


# --------------------------------------------------------------------------- #
def lire_arguments():
    """Ligne de commande > fichier de configuration > valeurs par défaut."""
    parser = argparse.ArgumentParser(
        description="Fusion des CSV de main_batch.py en un Excel, avec le nom des points de contrôle")
    # Positionnels (ordre : points, dossier, sortie) ou options nommées équivalentes
    parser.add_argument("pos", nargs="*", metavar="[points] [dossier_csv] [sortie]",
                        help="Arguments positionnels, dans cet ordre (facultatifs)")
    parser.add_argument("--points", help=f"Fichier de time series TELEMAC (défaut : {POINTS_FILE})")
    parser.add_argument("--dossier-csv", help=f"Dossier des CSV à fusionner (défaut : {INPUT_FOLDER})")
    parser.add_argument("--sortie", help=f"Classeur Excel produit (défaut : {OUTPUT_XLSX})")
    parser.add_argument("--dist-max", type=float,
                        help=f"Distance max point configuré / nœud retenu, en m (défaut : {DIST_MAX:g})")
    parser.add_argument("--config", help="Fichier de configuration TOML (voir config_exemple.toml)")
    args = parser.parse_args()
    if len(args.pos) > 3:
        parser.error("au plus trois arguments positionnels : points, dossier_csv, sortie")
    pos = args.pos + [None] * (3 - len(args.pos))
    args.points = args.points or pos[0]
    args.dossier_csv = args.dossier_csv or pos[1]
    args.sortie = args.sortie or pos[2]

    cfg = {}
    if args.config:
        with open(args.config, "rb") as f:
            cfg = tomllib.load(f)
    chemins = cfg.get("chemins", {})
    fusion = cfg.get("fusion", {})

    points_file  = args.points or fusion.get("fichier_points", POINTS_FILE)
    input_folder = args.dossier_csv or chemins.get("dossier_csv", INPUT_FOLDER)
    output_xlsx  = args.sortie or fusion.get("sortie_xlsx", OUTPUT_XLSX)
    dist_max     = args.dist_max if args.dist_max is not None else float(fusion.get("dist_max", DIST_MAX))
    return points_file, input_folder, output_xlsx, dist_max


def main():
    points_file, input_folder, output_xlsx, dist_max = lire_arguments()

    if not os.path.exists(points_file):
        print(f"❌ Fichier de points introuvable : {points_file}")
        sys.exit(1)

    points = lire_points(points_file)
    print(f"📍 {len(points)} points configurés lus dans {os.path.basename(points_file)}")

    all_data, all_columns = lire_csv(input_folder)
    if not all_data:
        print("❌ Aucune donnée lue.")
        sys.exit(1)
    coords = coords_colonnes(all_columns)

    non_associees = associer(points, coords, dist_max)

    # ---- Bilan console --------------------------------------------------- #
    manquants = [p for p in points if p["col"] is None]
    partages = [p for p in points if p["statut"].startswith("OK (")]
    print("\n" + "=" * 60)
    print(f"🔎 ASSOCIATION DES POINTS (distance max {dist_max:.0f} m)")
    print(f"  ✅ Trouvés   : {len(points) - len(manquants)}/{len(points)}")
    print(f"  ❌ Manquants : {len(manquants)}")
    for p in manquants:
        prox = (f"nœud le plus proche à {p['dist']:.1f} m ({p['col_proche']})"
                if p.get("col_proche") else "aucun nœud")
        print(f"     - [{p['id']}] {p['nom']}  (X={p['x']}, Y={p['y']}) → {prox}")
    if partages:
        print(f"  ⚠ Points partageant un même nœud :")
        for p in partages:
            print(f"     - [{p['id']}] {p['nom']} : {p['statut']}")
    if non_associees:
        print(f"  ⚠ {len(non_associees)} nœud(s) des CSV associé(s) à aucun point "
              f"(ajouté(s) en fin de tableau) :")
        for c in non_associees:
            print(f"     - {c}")
    print("=" * 60)

    print("💾 Export Excel...")
    exporter_xlsx(output_xlsx, points, all_data, coords, non_associees)
    print(f"✅ Terminé : {output_xlsx}")


if __name__ == "__main__":
    main()
